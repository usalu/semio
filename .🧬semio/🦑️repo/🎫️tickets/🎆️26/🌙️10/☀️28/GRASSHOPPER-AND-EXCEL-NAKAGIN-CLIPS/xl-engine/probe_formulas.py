"""🔬 Probe: how the `formulas` oracle handles the Nakagin workbook."""

import importlib.util
import sys
import time
from pathlib import Path

import formulas
import openpyxl
from openpyxl.workbook.defined_name import DefinedName

import re


def expand(formula, wb):
    def swap(match):
        names = wb.order()
        first, last = names.index(match[1]), names.index(match[2])
        return "SUM(" + ",".join("'" + n + "'!" + match[3] for n in names[first:last + 1]) + ")"
    return re.sub(r"SUM\('([^:']+):([^']+)'!(\$?[A-Z]+\$?\d+)\)", swap, formula)


HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "🗑️generated" / "xl-engine"
OUT.mkdir(parents=True, exist_ok=True)
MODEL = Path("C:/git/semio/🎓️teaching/🏛️architecture/📊️excel/🎬️clip/🧱️model/🐍️.py")
spec = importlib.util.spec_from_file_location("xlmodel", MODEL)
xl = importlib.util.module_from_spec(spec)
sys.modules["xlmodel"] = xl
spec.loader.exec_module(xl)

wb = xl.workbook()
book = openpyxl.Workbook()
book.remove(book.active)
for sheet in wb.sheets:
    ws = book.create_sheet(sheet.name)
    for key, cell in sheet.cells.items():
        ws[key] = expand(cell.formula, wb) if cell.formula else cell.value
for name, ref in wb.names.items():
    book.defined_names[name] = DefinedName(name, attr_text=ref)
path = OUT / "probe.xlsx"
book.save(path)
start = time.time()
model = formulas.ExcelModel().loads(str(path)).finish()
solution = model.calculate()
print("[DEBUG] seconds", time.time() - start, "keys", len(solution))
keys = list(solution)
print("[DEBUG]", keys[:5])
mine = xl.evaluate(wb)
bad = 0
for (sheet, key), value in mine.items():
    cell = wb.cell(sheet, key)
    if not cell.formula:
        continue
    ref = f"'[probe.xlsx]{sheet.upper()}'!{key}"
    got = solution.get(ref)
    if got is None:
        print("[DEBUG] missing", ref)
        bad += 1
        continue
    v = got.value[0, 0]
    print("[DEBUG]", ref, v, value) if bad < 12 and not (isinstance(value, float) and abs(float(v) - value) < 1e-9 * max(1, abs(value))) else None
print("[DEBUG] missing", bad)
