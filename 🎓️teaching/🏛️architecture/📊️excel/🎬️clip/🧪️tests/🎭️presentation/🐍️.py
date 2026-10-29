"""🎭 Validate every Excel clip scene discovered by folder and the workbook toolkit that draws them, in English and German."""
import ast
import copy
import importlib
import importlib.util
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = next(path for path in Path(__file__).resolve().parents if (path / "🧰️framework").is_dir())
sys.path.insert(0, str(ROOT))
CLIP = ROOT / "🎓️teaching/🏛️architecture/📊️excel/🎬️clip"
CASES = json.loads(Path(__file__).with_name("🔣️.cases.json").read_text(encoding="utf-8"))
SERIES = "🎓️teaching.🏛️architecture.📊️excel.🎬️clip"
ENGINE = (CLIP / "🧱️model/🐍️.py").is_file()
RESERVED = {"🧱️model", "🎨️visuals", "🧪️tests", "🧪️cases", "__pycache__"}


def discover() -> list[dict]:
    """🔎 Every topic folder of the series that carries a scene class with a NARRATION list, found by reading the source."""
    found = []
    for folder in sorted(path for path in CLIP.iterdir() if path.is_dir() and path.name not in RESERVED and (path / "🐍️.py").is_file()):
        tree = ast.parse((folder / "🐍️.py").read_text(encoding="utf-8"))
        for node in tree.body:
            if isinstance(node, ast.ClassDef) and any(isinstance(item, ast.Assign) and any(getattr(t, "id", "") == "NARRATION" for t in item.targets) for item in node.body):
                found.append({"folder": folder.name, "scene": node.name, "script": folder / "🐍️.py", "module": f"{SERIES}.{folder.name}.🐍️"})
    return found


def scene_class(entry):
    """🎬 Import one discovered scene class."""
    return getattr(importlib.import_module(entry["module"]), entry["scene"])


def step_keys(script: Path) -> list[str]:
    """🔑 Literal first arguments of every ``self.step(...)`` in a scene source."""
    keys = []
    for node in ast.walk(ast.parse(script.read_text(encoding="utf-8"))):
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) and node.func.attr == "step" and node.args and isinstance(node.args[0], ast.Constant):
            keys.append(node.args[0].value)
    return keys


def import_literals(script: Path) -> list[str]:
    """📦 Dotted module names a scene imports by literal ``import_module`` calls or ``import`` statements."""
    names = []
    for node in ast.walk(ast.parse(script.read_text(encoding="utf-8"))):
        if isinstance(node, ast.Call) and node.args and isinstance(node.args[0], ast.Constant) and isinstance(node.args[0].value, str):
            function = node.func
            if (function.attr if isinstance(function, ast.Attribute) else getattr(function, "id", "")) == "import_module":
                names.append(node.args[0].value)
        elif isinstance(node, ast.Import):
            names.extend(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module and not node.level:
            names.append(node.module)
    return names


class FakeWorkbook:
    """📗 The smallest workbook view the window needs, built from portable cases."""

    file_name = "Nakagin.xlsx"

    def __init__(self, spec: dict):
        self.sheets, self.cells, self.formats = spec["sheets"], dict(spec["cells"]), spec["formats"]
        self.cells[CASES["long_formula"]["cell"]] = CASES["long_formula"]["de"]

    def formula(self, sheet, addr):
        value = self.cells.get(addr) if sheet == self.sheets[1] else None
        return value if isinstance(value, str) and value.startswith("=") else None

    def value(self, sheet, addr):
        if sheet != self.sheets[1]:
            return None
        value = self.cells.get(addr)
        return self.value(sheet, "B1") * self.value(sheet, "B2") if isinstance(value, str) and value.startswith("=") else value

    def number_format(self, sheet, addr):
        return self.formats.get(addr, "General")

    def column_width(self, sheet, col):
        return 14.0


class SceneContract(unittest.TestCase):
    """🏛️ Every Excel clip scene obeys the narration, title and import contract; absent scenes are skipped until their folder exists."""

    def test_series_folders_are_registered(self):
        """🗂️ Every discovered scene folder is one of the planned topics."""
        planned = {case["folder"] for case in CASES["topics"]}
        for entry in discover():
            with self.subTest(folder=entry["folder"]):
                self.assertIn(entry["folder"], planned)

    def test_narration_title_and_steps(self):
        """📜 Match narration rows, localized titles, handlers and literal step keys of every scene."""
        for entry in discover():
            with self.subTest(scene=entry["scene"], folder=entry["folder"]):
                scene = scene_class(entry)
                self.assertTrue(callable(getattr(scene, "construct")))
                handlers = [name for name in scene.__dict__ if name.startswith("present_")]
                self.assertEqual(len(handlers), 1, handlers)
                rows = scene.NARRATION
                keys = [row[0] for row in rows]
                self.assertEqual(len(keys), len(set(keys)), "duplicate narration keys")
                self.assertTrue(CASES["clauses"]["min"] <= len(rows) <= CASES["clauses"]["max"], len(rows))
                for row in rows:
                    self.assertEqual(len(row), 3)
                    self.assertTrue(row[1].strip() and row[2].strip(), row[0])
                    self.assertNotEqual(row[1], row[2], row[0])
                self.assertTrue(all(scene.TITLE.get(language, "").strip() for language in CASES["languages"]))
                used = step_keys(entry["script"])
                self.assertEqual(set(used) - set(keys), set(), "step keys without narration")
                self.assertEqual(set(keys) - set(used), set(), "narration without step")

    def test_imports_resolve(self):
        """🔗 Every literal module a scene imports exists."""
        for entry in discover():
            for name in import_literals(entry["script"]):
                if name.startswith("🎓️teaching."):
                    with self.subTest(scene=entry["scene"], module=name):
                        self.assertIsNotNone(importlib.util.find_spec(name))

    def test_literal_dependency_closure(self):
        """🔐 Include the Excel widgets, the workbook widgets and the BIM drafting helpers in render provenance."""
        runtime = importlib.import_module("🎓️teaching.🏛️architecture.⚡️energy.🏗️authoring.🎬️clip.⚙️runtime.🐍️")
        teaching = ROOT / "🎓️teaching"
        required = {teaching / path for path in CASES["dependencies"]}
        for entry in discover():
            with self.subTest(scene=entry["scene"]):
                actual = set(runtime.source_dependencies(entry["script"]))
                self.assertEqual(sorted(str(p.relative_to(teaching)) for p in required - actual), [])

    def test_native_lifecycle(self):
        """🎞️ Exercise real animation setup, interpolation and Cairo in both languages for every scene."""
        import numpy as np
        from manim import tempconfig

        output = Path(os.environ.get("SEMIO_TEACHING_OUTPUT_ROOT") or tempfile.mkdtemp(prefix="excel-presentation-"))
        output.mkdir(parents=True, exist_ok=True)
        for entry in discover():
            for language in CASES["languages"]:
                with self.subTest(scene=entry["scene"], language=language), patch.dict(os.environ, {"SEMIO_TEACHING_LANGUAGE": language, "SEMIO_TEACHING_AUDIENCE": "spectator"}), \
                        tempconfig({"media_dir": str(output / "media"), "text_dir": str(output / "text"), "pixel_width": 480, "pixel_height": 270, "frame_rate": 4,
                                    "renderer": "cairo", "dry_run": True, "disable_caching": True}):
                    base, seen, played = scene_class(entry), [], set()

                    def play(scene, *animations, **kwargs):
                        compiled = scene.compile_animations(*animations, **kwargs)
                        scene.add_mobjects_from_animations(compiled)
                        scene.animations = compiled
                        scene.begin_animations()
                        for animation in compiled:
                            pending, visited = [animation], set()
                            while pending:
                                item = pending.pop()
                                if id(item) in visited:
                                    continue
                                visited.add(id(item))
                                played.add(type(item).__name__)
                                pending.extend(getattr(item, "animations", ()))
                        for alpha in (0.0, 0.5, 1.0):
                            for animation in compiled:
                                animation.update_mobjects(0.25)
                                animation.interpolate(alpha)
                            scene.update_mobjects(0.25)
                        for animation in compiled:
                            animation.finish()
                            animation.clean_up_from_scene(scene)
                        scene.update_mobjects(0)
                        scene.renderer.time += max((animation.run_time for animation in compiled), default=0)

                    def wait(scene, duration=1, **kwargs):
                        scene.update_mobjects(0.25)
                        scene.renderer.time += duration

                    def step(scene, key, *animations, **kwargs):
                        seen.append(key)
                        result = super(probe, scene).step(key, *animations, **kwargs)
                        caption = set(scene._caption.get_family())
                        body = [member for member in scene.mobjects if member not in caption]
                        self.assertTrue(any(len(member.points) for item in body for member in item.get_family()), key)
                        for member in scene.get_mobject_family_members():
                            self.assertTrue(np.isfinite(member.points).all(), key)
                        scene.camera.reset()
                        scene.camera.capture_mobjects(body)
                        self.assertTrue(np.any(scene.camera.pixel_array != scene.camera.pixel_array[0, 0]), key)
                        return result

                    probe = type("NativeExcelProbe", (base,), {"play": play, "wait": wait, "step": step})
                    scene = probe()
                    scene.construct()
                    self.assertEqual(seen, [row[0] for row in base.NARRATION])
                    self.assertTrue(set(CASES["animations"]) <= played, played)


class WorkbookToolkit(unittest.TestCase):
    """📗 The workbook widgets: portable address and reference cases checked against openpyxl, and a real window rendered in both languages."""

    @classmethod
    def setUpClass(cls):
        cls.book = importlib.import_module(f"{SERIES}.🎨️visuals.📗️workbook.🐍️")
        cls.visuals = importlib.import_module(f"{SERIES}.🎨️visuals.🐍️")

    def test_addresses_and_ranges(self):
        """📍 Column letters, cells and ranges match the portable cases and the openpyxl oracle."""
        from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

        for case in CASES["addresses"]:
            with self.subTest(a1=case["a1"]):
                self.assertEqual(self.book.split_address(case["a1"]), (case["column"], case["row"]))
                self.assertEqual(self.book.column_index(self.book.column_letters(case["column"])), case["column"])
                self.assertEqual(self.book.column_letters(case["column"]), get_column_letter(case["column"]))
                self.assertEqual(self.book.column_index(get_column_letter(case["column"])), column_index_from_string(get_column_letter(case["column"])))
        for case in CASES["ranges"]:
            with self.subTest(a1=case["a1"]):
                self.assertEqual(list(self.book.parse_range(case["a1"])), case["box"])
                if case["a1"].replace("$", "").split(":")[0] <= case["a1"].replace("$", "").split(":")[-1] and ":" in case["a1"]:
                    self.assertEqual(list(range_boundaries(case["a1"].replace("$", ""))), case["box"])

    def test_reference_tokens_and_colours(self):
        """🌈 Reference tokens are found in order, never inside strings or function names, and take Excel's colours in order."""
        palette = self.book.REFERENCE_COLORS
        for case in CASES["references"]:
            with self.subTest(formula=case["formula"]):
                found = [text for kind, text in self.book.reference_tokens(case["formula"]) if kind == "ref"]
                self.assertEqual(found, case["refs"])
                self.assertEqual("".join(text for _, text in self.book.reference_tokens(case["formula"])), case["formula"])
                colors = self.book.reference_colors(case["formula"])
                self.assertEqual(list(colors.values()), list(palette[:len(colors)]))
        self.assertEqual(list(palette[:4]), ["#2F6FDE", "#D93025", "#8E3FBF", "#1E9E4A"])

    def test_localized_strings(self):
        """🌐 Ribbon tab names, error labels and tooltips exist in both languages and match the cases."""
        data = self.book.DATA
        for language in CASES["languages"]:
            self.assertEqual([entry[language] for entry in data["ribbon"]], CASES["ribbon"][language])
        for code, names in CASES["errors"].items():
            for language in CASES["languages"]:
                self.assertEqual(self.book.error_name(code, language), names[language])
                self.assertTrue(data["errors"][code]["tip"][language].strip())
        schema = json.loads((Path(self.book.__file__).parent / "🧬️schema/🔣️.schema.json").read_text(encoding="utf-8"))
        self.assertTrue(set(schema["required"]) <= set(data))

    def test_punctuation_sits_on_the_baseline(self):
        """➖ A lone minus, asterisk or equals sign is placed by the text baseline, so the minus stays at x-height instead of dropping to an underscore."""
        capital = self.book._put("H", 0, 0, "l", 19.2)
        for sign in ("-", "*", "="):
            with self.subTest(sign=sign):
                shape = self.book._put(sign, 0, 0, "l", 19.2)
                self.assertGreater(shape.get_center()[1], capital.get_bottom()[1] + 0.25 * capital.height)
        group = self.book.formula_group("=B2-B3", size=19.2)
        minus = group[2]
        self.assertGreater(minus.get_center()[1], group[0].get_bottom()[1] + 0.25 * group[0].height)

    def test_formula_glyphs(self):
        """🔣 Render =SUM(B2:B4)-B5: parentheses are round (curved, not bracket stems), minus sits at x-height, tokens keep their order and colours."""
        import numpy as np

        for language in CASES["languages"]:
            with self.subTest(language=language), patch.dict(os.environ, {"SEMIO_TEACHING_LANGUAGE": language}):
                group = self.book.formula_group("=SUM(B2:B4)-B5", size=19.2)
                self.assertEqual([token for token, _, _ in group.refs], ["B2:B4", "B5"])
                opening, bracket = self.book._label("(", 19.2, "#000000"), self.book._label("[", 19.2, "#000000")
                for glyph, curved in ((opening, True), (bracket, False)):
                    points = np.array([m.point_from_proportion(a)[:2] for m in glyph.get_family() if len(m.points) for a in np.linspace(0, 1, 400)])
                    top, mid = points[points[:, 1] > points[:, 1].max() - 0.03 * np.ptp(points[:, 1])], points[abs(points[:, 1] - (points[:, 1].max() + points[:, 1].min()) / 2) < 0.2 * np.ptp(points[:, 1])]
                    bulge = (top[:, 0].min() - mid[:, 0].min()) / np.ptp(points[:, 0])
                    self.assertEqual(bulge > 0.2, curved, (glyph, bulge))
                minus = [m for m in group.submobjects if m not in [shape for _, _, shape in group.refs]][-1]
                self.assertGreater(minus.get_center()[1], group.submobjects[0].get_bottom()[1] + 0.25 * group.submobjects[0].height)

    def test_long_formula_wraps_inside_the_window(self):
        """📏 A 120-character German formula shrinks to 16.2, wraps into an expanded bar, stays inside the window and keeps every reference colour."""
        from manim import tempconfig

        case = CASES["long_formula"]
        self.assertGreaterEqual(len(case["de"]), 120)
        with patch.dict(os.environ, {"SEMIO_TEACHING_LANGUAGE": "de"}), tempconfig({"pixel_width": 480, "pixel_height": 270, "frame_rate": 1, "renderer": "cairo", "dry_run": True,
                                                                                       "disable_caching": True, "media_dir": tempfile.mkdtemp(prefix="excel-long-")}):
            probe = type("Probe", (self.visuals.ClipScene,), {"xl_localize": lambda scene, formula, lang=None: formula,
                                                                "xl_display": lambda scene, value, fmt="General", lang=None: f"{value:g}"})
            scene = probe()
            window = scene.xl_window(FakeWorkbook(CASES["window"]), "Type A", cols="A:C", rows=(1, 6), selected=case["cell"])
            field = window.formula_bar[1].current
            self.assertGreater(len(field.lines), 1)
            self.assertEqual("".join(field.lines), case["de"])
            left, right = window.x0, window.right
            for member in field.get_family():
                if len(member.points):
                    self.assertGreaterEqual(member.get_left()[0], left - 1e-6)
                    self.assertLessEqual(member.get_right()[0], right + 1e-6)
                    self.assertGreaterEqual(member.get_bottom()[1], window.bottom)
            window.edit(case["de"])
            self.assertEqual(len(window.frames.refresh().current), 1)

    def test_selection_tint_stays_under_text(self):
        """🟩 Draw order is fill, tint, gridlines, text, frames, border before and after a played selection and recalculation: the range tint never greys the values."""
        from manim import tempconfig

        model = importlib.import_module(f"{SERIES}.🧱️model.🐍️")
        book = model.workbook()
        after = model.apply(book, [{"type": "set-cell", "sheet": "Type A", "cell": "B3", "value": 5.0}])
        with patch.dict(os.environ, {"SEMIO_TEACHING_LANGUAGE": "en", "SEMIO_TEACHING_AUDIENCE": "spectator"}), tempconfig({
                "pixel_width": 480, "pixel_height": 270, "frame_rate": 4, "renderer": "cairo", "dry_run": True, "disable_caching": True,
                "media_dir": tempfile.mkdtemp(prefix="excel-order-")}):
            visuals = self.visuals
            from manim import FadeIn
            checks = []

            class Order(visuals.ClipScene):
                def construct(scene):
                    window = scene.xl_window(book, "Type A", rows=(1, 6), selected="B2")
                    scene.add(window, scene.xl_ref_frame(window, "B3:B4", color="#D93025"))
                    for moment in ("start", "select", "recalculate"):
                        if moment == "select":
                            scene.play(window.select("B3:B4"), run_time=0.1)
                        elif moment == "recalculate":
                            scene.play(window.recalculate(after), run_time=0.1)
                        family = list(scene.get_mobject_family_members())
                        position = lambda mobject: next(i for i, member in enumerate(family) if member is mobject)
                        text = position(window.cell("B3").text.get_family()[-1] if window.cell("B3").text.submobjects else window.cell("B3").text)
                        checks.append((moment, position(window.tint), position(window.gridlines), text, position(window.active), len(family)))
                    twin = scene.xl_window(after, "Type A", rows=(1, 6), selected="B2")
                    scene.play(FadeIn(twin), twin.select("B3:B4"), run_time=0.1)
                    hop = twin.select("B4")
                    hop.begin()
                    hop.interpolate(1.0)
                    hop.finish()
                    family = list(scene.get_mobject_family_members())
                    position = lambda mobject: next(i for i, member in enumerate(family) if member is mobject)
                    checks.append(("fade-in", position(twin.tint), position(twin.gridlines), position(twin.cell("B3").text.get_family()[-1]), position(twin.active), len(family)))


            Order().render()
        for moment, tint, grid, text, border, frame in checks:
            with self.subTest(moment=moment):
                self.assertLess(tint, grid)
                self.assertLess(grid, text)
                self.assertLess(text, border)

    def test_title_bar_keeps_the_file_name_readable(self):
        """🟢 In narrow windows the search pill shrinks to a magnifier and finally disappears; nothing ever sits under the file name."""
        from manim import tempconfig

        for language in CASES["languages"]:
            for width in (3.4, 4.2, 5.2, 6.8, 12.4):
                with self.subTest(language=language, width=width), patch.dict(os.environ, {"SEMIO_TEACHING_LANGUAGE": language}), tempconfig({
                        "pixel_width": 480, "pixel_height": 270, "frame_rate": 1, "renderer": "cairo", "dry_run": True, "disable_caching": True,
                        "media_dir": tempfile.mkdtemp(prefix="excel-title-")}):
                    probe = type("Probe", (self.visuals.ClipScene,), {"xl_localize": lambda scene, formula, lang=None: formula,
                                                                        "xl_display": lambda scene, value, fmt="General", lang=None: f"{value:g}"})
                    window = probe().xl_window(FakeWorkbook(CASES["window"]), "Type A", width=width, height=3.2, cols="A:B", rows=(1, 4))
                    members = window.title.submobjects
                    name = members[10]
                    others = [member for member in members[11:] if len(member.points)]
                    self.assertGreater(name.width, 0.5)
                    self.assertLessEqual(name.get_right()[0], window.right - 0.9)
                    for member in others:
                        self.assertGreaterEqual(member.get_left()[0], name.get_right()[0] + 0.2 - 1e-6, member)
                    if width >= 6.8:
                        self.assertGreaterEqual(len(others), 4)

    def test_window_renders_in_both_languages(self):
        """🪟 A window builds, answers cell lookups, follows a selection, an edit and a tracker, and paints green Excel chrome on Cairo."""
        import numpy as np
        from manim import ValueTracker, tempconfig

        spec = CASES["window"]
        book = FakeWorkbook(spec)
        for language in CASES["languages"]:
            with self.subTest(language=language), patch.dict(os.environ, {"SEMIO_TEACHING_LANGUAGE": language}), tempconfig({
                    "pixel_width": 480, "pixel_height": 270, "frame_rate": 1, "renderer": "cairo", "dry_run": True, "disable_caching": True,
                    "media_dir": tempfile.mkdtemp(prefix="excel-window-")}):
                probe = type("Probe", (self.visuals.ClipScene,), {
                    "xl_localize": lambda scene, formula, lang=None: formula,
                    "xl_display": lambda scene, value, fmt="General", lang=None: scene.number(value, 2) if fmt != "General" else f"{value:g}"})
                scene = probe()
                window = scene.xl_window(book, "Type A", selected=spec["selected"], cols="A:C", rows=(1, 6), editing=spec["editing"])
                self.assertEqual(set(window.tabs), set(book.sheets))
                self.assertEqual(window.active_address(), spec["selected"])
                self.assertEqual(len(window.range(spec["range"])), 2)
                for a1, code in (("C2", "#SPILL!"), ("C3", "#REF!")):
                    shown = self.book.error_name(code, language)
                    self.assertEqual(window.cell(a1).string, shown)
                    glyphs = window.cell(a1).text.get_family()[-1].parent if False else window.cell(a1).text
                    label = glyphs[-1] if len(glyphs) == 2 else glyphs
                    self.assertEqual(len(label.submobjects), len(shown.replace(" ", "")), "error code lost glyphs")
                    self.assertLessEqual(label.width, window.rect(a1)[2] - window.rect(a1)[0] + 0.7)
                left, top, right, bottom = window.rect("B1")
                self.assertTrue(left < right and top > bottom)
                self.assertLess(window.rect("A1")[2], window.rect("C1")[2])
                tracker = ValueTracker(spec["bound"]["tracker"])
                window.bind_value(spec["bound"]["cell"], tracker)
                self.assertEqual(len(window.frames.current), 2)
                animation = window.select(spec["range"])
                self.assertEqual(window.state["selected"], spec["range"])
                self.assertIsNotNone(animation)
                window.edit(None)
                self.assertEqual(len(window.frames.refresh().current), 0)
                scene.add(window, scene.xl_trace(window, "B1", "B3"), scene.xl_error(window, "B3", "#DIV/0!"), scene.xl_ref_frame(window, "B1:B2", color="#D93025", dashed=True),
                          scene.xl_fill_handle(window, "B3"), scene.xl_formula_text("=SUM(B1:B2,A1)", language=language), scene.xl_cell_pop(window, "B3", at=(0, 0, 0)),
                          scene.xl_bar_chart([("A", 8.7), ("B", tracker)], width=3, height=2))
                twin = copy.deepcopy(window)
                self.assertEqual(len(twin.cells), len(window.cells))
                self.assertIs(window.scene, scene)
                from manim import FadeIn, FadeOut
                for animation in (FadeIn(window), FadeOut(window)):
                    animation.begin()
                    animation.interpolate(0.5)
                    animation.finish()
                tracker.set_value(7.0)
                scene.update_mobjects(0.1)
                scene.camera.reset()
                scene.camera.capture_mobjects(scene.mobjects)
                pixels = scene.camera.pixel_array
                green = np.array([0x21, 0x73, 0x46])
                self.assertTrue(np.any(np.all(np.abs(pixels[..., :3].astype(int) - green) < 6, axis=-1)), "title bar green is missing")
                self.assertTrue(all(np.isfinite(member.points).all() for member in scene.get_mobject_family_members()))

    @unittest.skipUnless(ENGINE, "XL engine not built yet")
    def test_engine_hooks(self):
        """🧮 The default localize and display hooks come from the Excel engine."""
        scene = type("Probe", (importlib.import_module(f"{SERIES}.🎨️visuals.🐍️").ClipScene,), {})()
        with patch.dict(os.environ, {"SEMIO_TEACHING_LANGUAGE": "de"}):
            self.assertEqual(scene.xl_localize(CASES["engine"]["formula"]), CASES["engine"]["de"])


    @unittest.skipUnless(ENGINE, "XL engine not built yet")
    def test_window_over_the_nakagin_workbook(self):
        """📗 Every visible cell shows what the engine says it shows, in both languages, and a recalculation morphs exactly the changed cells."""
        from manim import tempconfig

        model = importlib.import_module(f"{SERIES}.🧱️model.🐍️")
        book = model.workbook()
        after = model.apply(book, [{"type": "set-cell", "sheet": "Type A", "cell": "B2", "value": 3.0}])
        for language in CASES["languages"]:
            with self.subTest(language=language), patch.dict(os.environ, {"SEMIO_TEACHING_LANGUAGE": language}), tempconfig({
                    "pixel_width": 480, "pixel_height": 270, "frame_rate": 1, "renderer": "cairo", "dry_run": True, "disable_caching": True,
                    "media_dir": tempfile.mkdtemp(prefix="excel-engine-")}):
                scene = type("Probe", (self.visuals.ClipScene,), {})()
                window = scene.xl_window(book, "Type A", cols="A:F", rows=(1, 10), selected="B9")
                values = model.evaluate(book)
                for a1, holder in window.cells.items():
                    self.assertEqual(holder.string or "", model.shown(book, values, "Type A", a1, language), a1)
                self.assertEqual(window.cell("B9").string, model.shown(book, values, "Type A", "B9", language))
                self.assertEqual(list(window.tabs), book.order())
                self.assertEqual(window.view.tab_title("Type A"), model.tab_name(book, "Type A", language))
                window.recalculate(after)
                changed = {a1 for a1, holder in window.cells.items() if holder.string != model.shown(book, values, "Type A", a1, language)}
                self.assertTrue({"B2", "D3", "D4"} <= changed)
                self.assertEqual(window.cell("B2").string, model.shown(after, model.evaluate(after), "Type A", "B2", language))
                self.assertEqual(window.cell("B5").string, model.shown(book, values, "Type A", "B5", language))


if __name__ == "__main__":
    unittest.main()
