"""🧪 Portable Excel teaching contracts with third-party oracles (openpyxl references and copy translation, formulas workbook evaluation)."""

import ast
import importlib.util
import io
import json
import math
import random
import re
import sys
import tempfile
import tokenize
import unittest
import warnings
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

warnings.simplefilter("ignore")
import openpyxl  # noqa: E402
import openpyxl.worksheet.table  # noqa: E402
from openpyxl.formula.translate import Translator, TranslatorError  # noqa: E402
from openpyxl.utils import cell as oxl_cell  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402

sys.dont_write_bytecode = True
CLIP = Path(__file__).resolve().parent.parent
CASES = json.loads((CLIP / "🧪️cases/🔣️.json").read_text(encoding="utf-8"))
ROOT = next(parent for parent in CLIP.parents if (parent / "🧰️framework").is_dir())
NAKAGIN = ROOT / "🎓️teaching/🏛️architecture/🏢️nakagin/🐍️.py"
TICKET = ROOT / ".🧬semio/🦑️repo/🎫️tickets/🎆️26/🌙️10/☀️28/GRASSHOPPER-AND-EXCEL-NAKAGIN-CLIPS"


def load_module(name, path):
    """🧱 Load a model file by path without importing drawing or speech engines."""
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


xl = load_module("architecture_excel_model", CLIP / "🧱️model/🐍️.py")


def key(text):
    """🔑 'Data!B1' to ('Data', 'B1')."""
    sheet, _, cell = text.rpartition("!")
    return sheet, cell


def build(sheets):
    """🏗️ A workbook from {sheet: {address: constant or "=formula"}}."""
    def make(item):
        if isinstance(item, str) and item.startswith("="):
            return xl.Cell(formula=item)
        return xl.Cell(value=float(item) if isinstance(item, int) and not isinstance(item, bool) else item)
    return xl.Workbook(tuple(xl.Sheet(name, {address: make(item) for address, item in cells.items()}, {}) for name, cells in sheets.items()), {})


def same(a, b):
    """⚖️ Value equality with 1e-9 relative tolerance, errors by code."""
    if isinstance(a, xl.XlError):
        return str(b) == a.code if not isinstance(b, dict) else b.get("error") == a.code
    if isinstance(b, dict):
        return False
    if isinstance(a, bool) or isinstance(b, bool):
        return bool(a) == bool(b) and isinstance(a, bool) == isinstance(b, bool)
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-12)
    return a == b


def expected_value(expected):
    """🎯 Portable expected value to engine value."""
    if isinstance(expected, dict):
        return xl.ERR[expected["error"]]
    return float(expected) if isinstance(expected, (int, float)) and not isinstance(expected, bool) else expected


def expand_structured(formula, table, host_row=None):
    """📐 Independent oracle: expand Table[Column], Table[@Column], Table[[#This Row],[Column]], Table[#Headers|#All|#Data] and Table[[A]:[B]] into absolute A1 ranges."""
    name = re.escape(table["name"])
    heads = [head.casefold() for head in table["headers"]]
    col0, row0, count = table["col"], table["row"], table["nrows"]
    letter = lambda index: oxl_cell.get_column_letter(col0 + heads.index(index.casefold()) if isinstance(index, str) else col0 + index)
    last = len(heads) - 1
    pattern = re.compile(
        r"(?:" + name + r")?\[\[#This Row\],\[(?P<tr>[^\]]+)\]\]"
        r"|(?:" + name + r")?\[@(?P<at>[^\]]+)\]"
        r"|" + name + r"\[\[(?P<s1>[^\]]+)\]:\[(?P<s2>[^\]]+)\]\]"
        r"|(?P<hd>" + name + r"\[#Headers\])"
        r"|(?P<al>" + name + r"\[#All\])"
        r"|(?P<da>" + name + r"\[#Data\])"
        r"|" + name + r"\[(?P<c>[^\[\]#@]+)\]")

    def swap(match):
        if match["tr"] or match["at"]:
            return letter(match["tr"] or match["at"]) + str(host_row)
        if match["s1"]:
            return f"${letter(match['s1'])}${row0 + 1}:${letter(match['s2'])}${row0 + count}"
        if match["hd"]:
            return f"${letter(0)}${row0}:${letter(last)}${row0}"
        if match["al"]:
            return f"${letter(0)}${row0}:${letter(last)}${row0 + count}"
        if match["da"]:
            return f"${letter(0)}${row0 + 1}:${letter(last)}${row0 + count}"
        return f"${letter(match['c'])}${row0 + 1}:${letter(match['c'])}${row0 + count}"
    return pattern.sub(swap, formula)


def program_table():
    """🗃️ The Program table of the authored workbook as plain data for the oracle."""
    spec = next(sheet["table"] for sheet in xl.load()["sheets"] if "table" in sheet)
    col, row = xl.split_address(spec["at"])
    return {"name": spec["name"], "col": col, "row": row, "headers": [c["header"]["en"] for c in spec["columns"]], "nrows": len(spec["rows"])}


def structured_scratch():
    """🧪 The scratch table of the cases with the oracle-flagged formulas in column F."""
    spec = CASES["tables"]
    wb = build({spec["sheet"]: {**spec["cells"], **{f"F{index + 10}": "=" + case["f"] for index, case in enumerate(spec["formulas"])}}})
    return xl.apply(wb, [{"type": "define-table", "name": spec["table"]["name"], "sheet": spec["sheet"], "ref": spec["table"]["ref"]}])


class ReferenceTests(unittest.TestCase):
    def test_columns_match_openpyxl(self):
        for letters, number in CASES["references"]["columns"]:
            self.assertEqual(xl.column_number(letters), number)
            self.assertEqual(xl.column_letters(number), letters)
            self.assertEqual(oxl_cell.column_index_from_string(letters), number)
            self.assertEqual(oxl_cell.get_column_letter(number), letters)
        rng = random.Random(7)
        for _ in range(500):
            number = rng.randint(1, xl.MAX_COL)
            self.assertEqual(xl.column_letters(number), oxl_cell.get_column_letter(number))
            self.assertEqual(xl.column_number(xl.column_letters(number)), number)

    def test_parse_and_format(self):
        for case in CASES["references"]["parse"]:
            ref = xl.parse_ref(case["text"])
            self.assertEqual(xl.format_ref(ref), case["canonical"], case["text"])
            self.assertEqual(xl.format_ref(xl.parse_ref(case["canonical"])), case["canonical"])

    def test_parse_matches_openpyxl_boundaries(self):
        for case in CASES["references"]["parse"]:
            if "!" in case["text"]:
                continue
            ref = xl.parse_ref(case["text"])
            expected = oxl_cell.range_boundaries(case["text"].replace("$", ""))
            if isinstance(ref, xl.Ref):
                got = (ref.col, ref.row, ref.col, ref.row)
            else:
                got = (ref.start.col, ref.start.row, ref.end.col, ref.end.row)
            self.assertEqual(got, tuple(expected), case["text"])

    def test_single_cell_matches_openpyxl(self):
        rng = random.Random(11)
        for _ in range(300):
            col, row = rng.randint(1, xl.MAX_COL), rng.randint(1, xl.MAX_ROW)
            text = oxl_cell.get_column_letter(col) + str(row)
            ref = xl.parse_ref(text)
            self.assertEqual((oxl_cell.column_index_from_string(oxl_cell.coordinate_from_string(text)[0]), oxl_cell.coordinate_from_string(text)[1]), (ref.col, ref.row))
            self.assertEqual(xl.split_address(text), (col, row))
            self.assertEqual(xl.address(col, row), text)

    def test_invalid_references(self):
        for text in CASES["references"]["invalid"]:
            with self.assertRaises(xl.FormulaError, msg=text):
                xl.parse_ref(text)

    def test_sheet_quoting(self):
        for name, quoted in CASES["references"]["quote"]:
            self.assertEqual(xl.quote_sheet(name), quoted, name)
        for name in ("Type A", "Energy-Total"):
            self.assertEqual(xl.quote_sheet(name), openpyxl.utils.quote_sheetname(name), name)

    def test_f4_cycle(self):
        for before, after in CASES["references"]["cycle"]:
            self.assertEqual(xl.cycle_reference(before), after, before)
        text = "B2"
        seen = []
        for _ in range(4):
            text = xl.cycle_reference(text)
            seen.append(text)
        self.assertEqual(seen, ["$B$2", "B$2", "$B2", "B2"])

    def test_reference_spans_locate_text(self):
        formula = "=COUNTIFS($I$2:$I$141,$A11,'Type A'!M2:M141,B$10)"
        spans = xl.reference_spans(formula)
        self.assertEqual([formula[s:e] for s, e, _ in spans], ["$I$2:$I$141", "$A11", "'Type A'!M2:M141", "B$10"])
        self.assertEqual(len(xl.references(formula)), 4)


class TranslateTests(unittest.TestCase):
    def test_cases(self):
        for case in CASES["translate"]:
            self.assertEqual(xl.translate(case["formula"], case["from"], case["to"]), case["expected"], case)

    def test_openpyxl_oracle(self):
        checked = 0
        for case in CASES["translate"]:
            if case.get("oracle"):
                self.assertEqual(Translator(case["formula"], origin=case["from"]).translate_formula(case["to"]), case["expected"], case)
                checked += 1
        self.assertGreater(checked, 10)

    def test_random_copies_match_openpyxl(self):
        rng = random.Random(3)
        pieces = ["{c}{r}", "${c}{r}", "{c}${r}", "${c}${r}"]
        for _ in range(400):
            cells = []
            for _ in range(rng.randint(1, 3)):
                template = rng.choice(pieces)
                cells.append(template.format(c=xl.column_letters(rng.randint(5, 40)), r=rng.randint(5, 60)))
            formula = "=" + rng.choice(["{}", "SUM({}:{})", "{}+{}*2", "IF({}>{},{},0)"]).format(*cells, *cells, *cells)
            src, dst = xl.address(rng.randint(5, 40), rng.randint(5, 60)), xl.address(rng.randint(5, 40), rng.randint(5, 60))
            try:
                expected = Translator(formula, origin=src).translate_formula(dst)
            except TranslatorError:
                self.assertIn("#REF!", xl.translate(formula, src, dst), (formula, src, dst))
                continue
            self.assertEqual(xl.translate(formula, src, dst), expected, (formula, src, dst))

    def test_fill_is_inverse_of_translate(self):
        formula = "=SUM($B3:C$5)+D7"
        there = xl.translate(formula, "A1", "F9")
        self.assertEqual(xl.translate(there, "F9", "A1"), formula)


class ParseTests(unittest.TestCase):
    def test_trees(self):
        for case in CASES["parse"]:
            self.assertEqual(xl.parse(case["formula"]).data(), case["tree"], case["formula"])

    def test_invalid_formulas(self):
        for formula in CASES["parse_invalid"]:
            with self.assertRaises(xl.FormulaError, msg=formula):
                xl.parse(formula)

    def test_leading_equals_is_optional(self):
        self.assertEqual(xl.parse("1+2").data(), xl.parse("=1+2").data())

    def test_german_parse_reads_decimal_comma(self):
        self.assertEqual(xl.parse("=SUMME(1,5;2)", "de").data(), ["call", "SUM", [["num", 1.5], ["num", 2]]])


class EvaluateTests(unittest.TestCase):
    def test_scalar_cases(self):
        shared = CASES["evaluate"]["shared"]
        for case in CASES["evaluate"]["scalar"]:
            self.assertTrue(same(xl.calculate(case["f"], shared), expected_value(case["v"])), (case["f"], xl.calculate(case["f"], shared)))

    def test_cell_cases(self):
        for case in CASES["evaluate"]["cells"]:
            self.assertTrue(same(xl.calculate(case["f"], case["cells"]), expected_value(case["v"])), (case["f"], xl.calculate(case["f"], case["cells"])))

    def test_spill_and_blocked_spill(self):
        for case in CASES["evaluate"]["spill"]:
            wb = build({"S": {**case["cells"], case["at"]: case["formula"]}})
            values = xl.evaluate(wb)
            for address, expected in case["expected"].items():
                self.assertTrue(same(values[("S", address)], expected_value(expected)), (case, address))

    def test_error_propagation(self):
        case = CASES["errors"]["propagation"]
        values = xl.evaluate(build({"S": case["cells"]}))
        for address, expected in case["expected"].items():
            self.assertTrue(same(values[("S", address)], expected_value(expected)), address)

    def test_errors_are_values_not_exceptions(self):
        self.assertEqual(xl.calculate("1/0"), xl.ERR["#DIV/0!"])
        self.assertEqual(xl.value_data(xl.calculate("1/0")), {"error": "#DIV/0!"})

    def test_german_text_uses_german_names_only(self):
        for case in CASES["de_unknown"]:
            got = xl.calculate(case["de"], {}, "de")
            self.assertTrue(same(got, expected_value(case["error"] and {"error": case["error"]} if "error" in case else case["value"])), case)

    def test_rounding_is_half_away_from_zero_on_decimals(self):
        for number, digits in [(2.5, 0), (-2.5, 0), (1.005, 2), (0.285, 2), (1234.5678, -2), (2.675, 2)]:
            expected = float(Decimal(repr(number)).quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP)) if number > 0 else -float(Decimal(repr(-number)).quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP))
            self.assertEqual(xl.round_half_up(number, digits), expected)

    def test_formulas_oracle_on_flagged_cases(self):
        book = openpyxl.Workbook()
        book.remove(book.active)
        flagged = []
        for index, case in enumerate(CASES["evaluate"]["scalar"]):
            if case.get("oracle"):
                flagged.append((f"s{index}", case["f"], CASES["evaluate"]["shared"], case["v"]))
        for index, case in enumerate(CASES["evaluate"]["cells"]):
            if case.get("oracle"):
                flagged.append((f"c{index}", case["f"], case["cells"], case["v"]))
        for name, formula, cells, _ in flagged:
            sheet = book.create_sheet(name)
            for address, item in cells.items():
                sheet[address] = item
            sheet["Z99"] = "=" + formula
        directory = Path(tempfile.mkdtemp(prefix="xl-engine-"))
        path = directory / "cases.xlsx"
        book.save(path)
        import formulas
        solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
        for name, formula, cells, expected in flagged:
            got = solution[f"'[cases.xlsx]{name.upper()}'!Z99"].value[0, 0]
            want = expected_value(expected)
            if isinstance(want, xl.XlError):
                self.assertEqual(str(got), want.code, formula)
            elif isinstance(want, bool):
                self.assertEqual(bool(got), want, formula)
            elif isinstance(want, float):
                self.assertTrue(math.isclose(float(got), want, rel_tol=1e-9, abs_tol=1e-12), (formula, got, want))
            else:
                self.assertEqual(str(got), want, formula)
        self.assertGreater(len(flagged), 60)


class GraphTests(unittest.TestCase):
    def test_precedents_and_dependents(self):
        graph = CASES["graph"]
        wb = build({"S": graph["cells"]})
        for cell, expected in graph["precedents"].items():
            self.assertEqual([c for _, c in xl.precedents(wb, ("S", cell))], expected, cell)
        for cell, expected in graph["dependents"].items():
            self.assertEqual([c for _, c in xl.dependents(wb, ("S", cell))], expected, cell)

    def test_recalculation_order_is_topological(self):
        wb = xl.workbook()
        order = xl.recalculation_order(wb)
        position = {item: index for index, item in enumerate(order)}
        self.assertEqual(len(order), len(set(order)))
        for item in order:
            for other in xl.precedents(wb, item):
                if other in position:
                    self.assertLess(position[other], position[item], (other, item))

    def test_circular_references(self):
        case = CASES["graph"]["circular"]
        wb = build({"S": case["cells"]})
        self.assertEqual(sorted(c for _, c in xl.circular_cells(wb)), sorted(case["cycle"]))
        values = xl.evaluate(wb)
        for address, expected in case["values"].items():
            self.assertEqual(values[("S", address)], expected)

    def test_names_are_dependencies(self):
        wb = xl.workbook()
        self.assertIn(("Climate", "B2"), xl.precedents(wb, ("Type A", "B13")))
        self.assertIn(("Energy", "B13"), xl.precedents(wb, ("Type A", "E5")))


class LocalizeTests(unittest.TestCase):
    def test_localize_and_back(self):
        for case in CASES["localize"]:
            self.assertEqual(xl.localize(case["en"], "de"), case["de"], case["en"])
            self.assertEqual(xl.canonicalize(case["de"], "de"), case["en"], case["de"])
            self.assertEqual(xl.localize(case["en"], "en"), case["en"])

    def test_german_names_cover_every_function(self):
        self.assertEqual(set(xl.DE_FUNCTIONS), set(xl.FUNCTIONS) | set(xl.LAZY))

    def test_english_names_do_not_evaluate_in_german(self):
        self.assertEqual(xl.calculate("=SUM(1;2)", {}, "de"), xl.ERR["#NAME?"])
        self.assertEqual(xl.calculate("=SUMME(1;2)", {}, "de"), 3.0)

    def test_error_labels(self):
        for code, german in CASES["errors"]["labels"]:
            self.assertEqual(xl.localize_error(code, "de"), german)
            self.assertEqual(xl.localize_error(code, "en"), code)
            self.assertEqual(xl.display(xl.ERR[code], None, "de"), german)

    def test_evaluation_is_language_independent(self):
        values = xl.evaluate(xl.workbook())
        for sheet in xl.workbook().sheets:
            for address, cell in sheet.cells.items():
                if cell.formula:
                    german = xl.canonicalize(xl.localize(cell.formula, "de", xl.workbook()), "de", xl.workbook())
                    self.assertEqual(german, cell.formula)
        self.assertGreater(len(values), 1000)


class TextConstantTests(unittest.TestCase):
    def test_every_text_constant_is_a_bilingual_label(self):
        data = xl.load()
        for spec in data["sheets"]:
            for address, cell in spec["cells"].items():
                self.assertFalse(isinstance(cell.get("value"), str), (spec["name"], address))
            for column in spec.get("table", {}).get("columns", []):
                index = [c["key"] for c in spec["table"]["columns"]].index(column["key"])
                if "formula" in column:
                    continue
                strings = {row[index - sum(1 for c in spec["table"]["columns"][:index] if "formula" in c)] for row in spec["table"]["rows"]}
                if any(isinstance(item, str) for item in strings):
                    self.assertEqual(strings, set(column["labels"]), column["key"])

    def test_german_workbook_matches_in_german_texts(self):
        en, de = xl.workbook(), xl.workbook(language="de")
        ven, vde = xl.evaluate(en), xl.evaluate(de)
        for target in ("A4", "A7", "B10", "E10", "A15", "M2", "M141", "N2"):
            self.assertEqual(vde[("Program", target)] if target != "N2" else 0, ({"A4": "Typ A", "A7": "Typ D", "B10": "Typ A", "E10": "Typ D", "A15": "Typ C", "M2": "Typ A", "M141": vde[("Program", "M141")], "N2": 0})[target])
        for sheet in en.sheets:
            for address, cell in sheet.cells.items():
                if cell.formula and not isinstance(ven[(sheet.name, address)], str):
                    self.assertTrue(same(ven[(sheet.name, address)], vde[(sheet.name, address)]), (sheet.name, address))
        self.assertTrue(vde[("Program", "H2")].startswith("A-01-"))
        self.assertIn("O", {vde[("Program", f"K{row}")] for row in range(2, 142)})
        self.assertNotIn("E", {vde[("Program", f"K{row}")] for row in range(2, 142)})
        self.assertEqual([vde[("Program", f"A{row}")] for row in range(4, 8)], ["Typ A", "Typ B", "Typ C", "Typ D"])
        self.assertEqual([ven[("Program", f"A{row}")] for row in range(4, 8)], ["Type A", "Type B", "Type C", "Type D"])
        self.assertEqual(xl.shown(de, vde, "Program", "A4", "de"), "Typ A")

    def test_lookup_and_criteria_follow_the_language(self):
        for language, name in (("en", "Type C"), ("de", "Typ C")):
            wb = xl.workbook(language=language)
            values = xl.evaluate(wb)
            self.assertEqual(values[("Program", "A15")], name)
            self.assertEqual(values[("Program", "C15")], 34.0)
            self.assertEqual(values[("Program", "B6")], 34.0)
            self.assertEqual(values[("Program", "B8")], 140.0)


class StructuredReferenceTests(unittest.TestCase):
    def test_scratch_table_formulas(self):
        spec = CASES["tables"]
        wb = structured_scratch()
        values = xl.evaluate(wb)
        for address, expected in spec["rows"].items():
            self.assertEqual(values[(spec["sheet"], address)], expected, address)
        for index, case in enumerate(spec["formulas"]):
            got = values[(spec["sheet"], f"F{index + 10}")]
            self.assertTrue(same(got, expected_value(case["v"])), (case["f"], got))

    def test_formulas_oracle_on_expanded_structured_references(self):
        spec = CASES["tables"]
        wb = structured_scratch()
        mine = xl.evaluate(wb)
        table = {"name": "Capsules", "col": 1, "row": 1, "headers": ["Type", "Area", "Tower", "Calc"], "nrows": 3}
        book = openpyxl.Workbook()
        ws = book.active
        ws.title = spec["sheet"]
        for address, item in spec["cells"].items():
            ws[address] = expand_structured(item, table, xl.split_address(address)[1]) if isinstance(item, str) and item.startswith("=") else item
        flagged = [(index, case) for index, case in enumerate(spec["formulas"]) if case.get("oracle")]
        for index, case in flagged:
            ws[f"F{index + 10}"] = "=" + expand_structured(case["f"], table, index + 10)
        directory = Path(tempfile.mkdtemp(prefix="xl-structured-"))
        path = directory / "structured.xlsx"
        book.save(path)
        import formulas
        solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
        for address in ("D2", "D3", "D4"):
            self.assertEqual(float(solution[f"'[structured.xlsx]S'!{address}"].value[0, 0]), mine[("S", address)])
        for index, case in flagged:
            got = solution[f"'[structured.xlsx]S'!F{index + 10}"].value[0, 0]
            want = mine[("S", f"F{index + 10}")]
            if isinstance(want, str):
                self.assertEqual(str(got), want, case["f"])
            else:
                self.assertAlmostEqual(float(got), want, msg=case["f"])
        self.assertGreater(len(flagged), 12)

    def test_resolution_to_a1_ranges(self):
        wb = structured_scratch()
        for text, host, expected in (("Capsules[Type]", None, "S!A2:A4"), ("Capsules[#Headers]", None, "S!A1:D1"), ("Capsules[#All]", None, "S!A1:D4"), ("Capsules[[Area]:[Calc]]", None, "S!B2:D4"), ("Capsules[@Area]", "D3", "S!B3:B3"), ("[@Area]", "D4", "S!B4:B4"), ("Capsules[[#This Row],[Tower]]", "D2", "S!C2:C2")):
            rng = xl.structured_range(wb, text, "S", host)
            self.assertEqual(xl.format_ref(rng), expected, text)

    def test_dependency_graph_sees_through_structured_references(self):
        spec = CASES["tables"]
        wb = structured_scratch()
        for cell, expected in spec["dependencies"]["precedents"].items():
            self.assertEqual([c for _, c in xl.precedents(wb, ("S", cell))], expected, cell)
        index = next(i for i, case in enumerate(spec["formulas"]) if case["f"] == spec["dependencies"]["formula_precedents"]["f"])
        self.assertEqual([c for _, c in xl.precedents(wb, ("S", f"F{index + 10}"))], spec["dependencies"]["formula_precedents"]["cells"])
        for cell, expected in spec["dependencies"]["dependents"].items():
            self.assertEqual([c for _, c in xl.dependents(wb, ("S", cell)) if not c.startswith("F")], expected, cell)
        self.assertIn(("S", "F11"), xl.dependents(wb, ("S", "B2")))
        order = xl.recalculation_order(wb)
        self.assertLess(order.index(("S", "D2")), order.index(("S", f"F{index + 10}")))

    def test_reference_spans_resolve_structured_text(self):
        spec = CASES["tables"]["spans"]
        wb = structured_scratch()
        spans = xl.reference_spans(spec["formula"], wb, "S", spec["host"])
        self.assertEqual([spec["formula"][a:b] for a, b, _ in spans], spec["texts"])
        self.assertEqual([xl.format_ref(item) for _, _, item in spans], spec["ranges"])
        self.assertEqual(len(xl.reference_spans(spec["formula"])), 0)

    def test_table_events(self):
        spec = CASES["tables"]
        base = structured_scratch()
        for case in spec["events"]:
            wb = xl.apply(base, case["events"])
            values = xl.evaluate(wb)
            for target, expected in case["values"].items():
                self.assertTrue(same(values[key(target)], expected_value(expected)), (case["id"], target, values.get(key(target))))
            self.assertEqual({name: {"row": t.row, "col": t.col, "nrows": t.nrows} for name, t in wb.tables.items()}, case["tables"], case["id"])
            if "headers" in case:
                self.assertEqual(len(wb.tables["Capsules"].headers), case["headers"])

    def test_define_table_reads_headers_from_cells(self):
        wb = xl.apply(build({"S": {"A1": "X", "B1": "Y", "A2": 1, "B2": 2, "A3": 3, "B3": 4, "D1": "=SUM(T[Y])"}}), [{"type": "define-table", "name": "T", "sheet": "S", "ref": "A1:B3"}])
        self.assertEqual(xl.evaluate(wb)[("S", "D1")], 6.0)
        self.assertEqual([h["en"] for h in wb.tables["T"].headers], ["X", "Y"])

    def test_translate_leaves_structured_references_untouched(self):
        self.assertEqual(xl.translate("=XLOOKUP([@Type],$A$4:$A$7,$C$4:$C$7)+Capsules[Net area]", "N2", "N3"), "=XLOOKUP([@Type],$A$4:$A$7,$C$4:$C$7)+Capsules[Net area]")

    def test_german_text_with_structured_references_evaluates(self):
        wb = xl.workbook()
        text = xl.localize("=SUM(Capsules[Net area])+ROWS(Capsules[#Headers])", "de", wb)
        self.assertEqual(text, "=SUMME(Kapseln[Nettogrundfläche])+ZEILEN(Kapseln[#Kopfzeilen])")
        added = xl.apply(wb, [{"type": "set-cell", "sheet": "Energy", "cell": "C17", "formula": xl.canonicalize(text, "de", wb)}])
        self.assertAlmostEqual(xl.evaluate(added)[("Energy", "C17")], 140 * 8.74 + 1)

    def test_workbook_uses_the_neutral_table_name_in_both_languages(self):
        wb = xl.workbook()
        self.assertEqual(list(wb.tables), ["Capsules"])
        self.assertEqual((wb.tables["Capsules"].sheet, wb.tables["Capsules"].nrows), ("Program", 140))
        self.assertEqual(xl.table_headers("de")["Net area"], "Nettogrundfläche")
        self.assertEqual([h["de"] for h in wb.tables["Capsules"].headers], ["Kapsel", "Turm", "Geschoss", "Fassade", "Platz", "Typ", "Nettogrundfläche"])
        self.assertIn("Capsules[Type]", wb.cell("Program", "B4").formula)
        self.assertEqual(xl.formula_text(wb, "Program", "B4", "de"), "=ZÄHLENWENN(Kapseln[Typ];A4)")
        self.assertEqual(xl.table_titles("de"), {"Capsules": "Kapseln"})
        self.assertEqual(xl.table_titles("en", wb), {"Capsules": "Capsules"})
        self.assertIs(wb.table("Kapseln"), wb.table("Capsules"))
        self.assertEqual(CASES["workbook"]["tables"]["Capsules"]["de"], wb.tables["Capsules"].labels["de"])

    def test_german_names_evaluate_after_canonicalize_and_directly(self):
        wb = xl.workbook()
        german = "=SUMME(Kapseln[Nettogrundfläche])+ZEILEN(Kapseln)+ZEILEN(Kapseln[[#Kopfzeilen];[Typ]])"
        english = xl.canonicalize(german, "de", wb)
        self.assertEqual(english, "=SUM(Capsules[Net area])+ROWS(Capsules)+ROWS(Capsules[[#Headers],[Type]])")
        self.assertEqual(xl.localize(english, "de", wb), german)
        rng = xl.structured_range(wb, "Kapseln[Nettogrundfläche]")
        self.assertEqual(xl.format_ref(rng), "Program!N2:N141")
        self.assertEqual(xl.format_ref(xl.structured_range(wb, "Capsules[Net area]")), "Program!N2:N141")

    def test_precedents_of_the_structured_program_formulas(self):
        wb = xl.workbook()
        self.assertEqual(set(xl.precedents(wb, ("Program", "B4"))), {("Program", "A4")} | {("Program", f"M{row}") for row in range(2, 142)})
        self.assertEqual(xl.precedents(wb, ("Program", "H2")), [("Program", c + "2") for c in "IJKL"])


class TabNameTests(unittest.TestCase):
    def test_tabs_are_user_data_localized_per_language(self):
        wb = xl.workbook()
        for sheet, texts in CASES["workbook"]["tabs"].items():
            for language in ("en", "de"):
                self.assertEqual(xl.tab_name(wb, sheet, language), texts[language])
        self.assertEqual(xl.sheet_titles("de")["Type A"], "Typ A")
        self.assertEqual(xl.tab_name(build({"Other": {}}), "Other", "de"), "Other")

    def test_every_workbook_formula_round_trips_through_german(self):
        wb = xl.workbook()
        for sheet in wb.sheets:
            for address, cell in sheet.cells.items():
                if not cell.formula:
                    continue
                german = xl.localize(cell.formula, "de", wb)
                for name in wb.order():
                    if name != xl.tab_name(wb, name, "de"):
                        self.assertNotRegex(german, r"(?<![\w'])" + re.escape(name) + r"[!:]|'" + re.escape(name) + r"[:']", (sheet.name, address, german))
                self.assertEqual(xl.canonicalize(german, "de", wb), cell.formula, (sheet.name, address))

    def test_german_formula_evaluates_to_the_same_value(self):
        wb = xl.workbook()
        values = xl.evaluate(wb)
        german = xl.localize("=SUM('Type A:Type D'!B13)+Programm!B4", "de", wb)
        self.assertEqual(german, "=SUMME('Typ A:Typ D'!B13)+Programm!B4")
        text = xl.canonicalize(german, "de", wb)
        self.assertEqual(text, "=SUM('Type A:Type D'!B13)+Programm!B4".replace("Programm", "Program"))
        added = xl.apply(wb, [{"type": "set-cell", "sheet": "Energy", "cell": "C16", "formula": text}])
        self.assertAlmostEqual(xl.evaluate(added)[("Energy", "C16")], values[("Energy", "B16")] + values[("Program", "B4")])

    def test_renamed_sheet_loses_its_tab_label(self):
        wb = xl.apply(xl.workbook(), [{"type": "rename-sheet", "sheet": "Climate", "name": "Weather"}])
        self.assertEqual(xl.tab_name(wb, "Weather", "de"), "Weather")
        self.assertEqual(xl.tab_name(wb, "Energy", "de"), "Energie")


class DisplayTests(unittest.TestCase):
    def test_cases(self):
        for case in CASES["display"]:
            value = xl.ERR[case["error"]] if "error" in case else case["value"]
            for language in ("en", "de"):
                self.assertEqual(xl.display(value, case["format"], language), case[language], (case, language))

    def test_decimal_oracle(self):
        rng = random.Random(5)
        for _ in range(400):
            value = round(rng.uniform(-50000, 50000), rng.randint(0, 6))
            decimals = rng.randint(0, 4)
            grouped = rng.random() < 0.5
            pattern = ("#,##0" if grouped else "0") + ("." + "0" * decimals if decimals else "")
            exact = Decimal(repr(abs(value))).quantize(Decimal(1).scaleb(-decimals), rounding=ROUND_HALF_UP)
            text = format(exact, ",f" if grouped else "f")
            sign = "-" if value < 0 else ""
            self.assertEqual(xl.display(value, pattern, "en"), sign + text, (value, pattern))
            self.assertEqual(xl.display(value, pattern, "de"), sign + text.translate(str.maketrans(",.", ".,")), (value, pattern))

    def test_percent_oracle(self):
        rng = random.Random(9)
        for _ in range(200):
            value = round(rng.uniform(0, 3), 4)
            exact = (Decimal(repr(value)) * 100).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
            self.assertEqual(xl.display(value, "0.0%", "en"), f"{exact}%")

    def test_numbers_from_one_thousand_show_a_thousands_separator(self):
        wb = xl.workbook()
        values = xl.evaluate(wb)
        checked = 0
        for (sheet, address), value in values.items():
            if isinstance(value, float) and abs(value) >= 1000 and (wb.cell(sheet, address).formula or wb.cell(sheet, address).value is not None):
                fmt = wb.cell(sheet, address).format
                self.assertIn("#,##0", fmt or "", (sheet, address, fmt))
                self.assertIn(",", xl.display(value, fmt, "en"))
                self.assertIn(".", xl.display(value, fmt, "de").replace(",", ""))
                checked += 1
        self.assertGreater(checked, 10)

    def test_stored_value_is_unchanged_by_format(self):
        wb = xl.workbook()
        values = xl.evaluate(wb)
        self.assertAlmostEqual(values[("Type A", "B2")], 2.5)
        self.assertEqual(xl.display(values[("Type A", "B2")], wb.cell("Type A", "B2").format, "en"), "2.50 m")


class EventTests(unittest.TestCase):
    def check(self, wb, case):
        for target, formula in case.get("formulas", {}).items():
            cell = wb.cell(*key(target))
            self.assertIsNotNone(cell, (case["id"], target))
            self.assertEqual(cell.formula, formula, (case["id"], target))
        for target in case.get("absent", []):
            self.assertIsNone(wb.cell(*key(target)), (case["id"], target))
        for target in case.get("constants", []):
            self.assertIsNone(wb.cell(*key(target)).formula, (case["id"], target))
        if "order" in case:
            self.assertEqual(wb.order(), case["order"], case["id"])
        values = xl.evaluate(wb)
        for target, expected in case.get("values", {}).items():
            self.assertTrue(same(values[key(target)], expected_value(expected)), (case["id"], target, values[key(target)]))

    def test_event_cases(self):
        base = build(CASES["events"]["base"])
        before = xl.evaluate(base)
        for case in CASES["events"]["cases"]:
            self.check(xl.apply(base, case["events"]), case)
        self.assertEqual(xl.evaluate(base), before)

    def test_three_d_reference_cases(self):
        base = build(CASES["events"]["three_d"]["base"])
        for case in CASES["events"]["three_d"]["cases"]:
            self.check(xl.apply(base, case["events"]), case)

    def test_apply_is_a_fold(self):
        base = build(CASES["events"]["base"])
        events = [{"type": "set-cell", "sheet": "Data", "cell": "A1", "value": 7}, {"type": "insert-row", "sheet": "Data", "index": 1}, {"type": "rename-sheet", "sheet": "Other", "name": "O2"}]
        one = xl.apply(base, events)
        step = base
        for event in events:
            step = xl.apply(step, [event])
        self.assertEqual({s.name: {k: (c.value, c.formula) for k, c in s.cells.items()} for s in one.sheets}, {s.name: {k: (c.value, c.formula) for k, c in s.cells.items()} for s in step.sheets})

    def test_unknown_event_is_rejected(self):
        with self.assertRaises(KeyError):
            xl.apply(build({"S": {}}), [{"type": "explode"}])

    def test_workbook_edit_keeps_dependents_current(self):
        wb = xl.workbook()
        base = xl.evaluate(wb)
        better = xl.apply(wb, [{"type": "set-cell", "sheet": "Energy", "cell": "B14", "value": 1.0}])
        after = xl.evaluate(better)
        self.assertLess(after[("Energy", "D8")], base[("Energy", "D8")])
        self.assertLess(after[("Type A", "B14")], base[("Type A", "B14")])
        self.assertEqual(after[("Program", "B8")], base[("Program", "B8")])

    def test_deleting_a_type_sheet_breaks_program_with_ref_errors(self):
        wb = xl.apply(xl.workbook(), [{"type": "delete-sheet", "sheet": "Type B"}])
        values = xl.evaluate(wb)
        self.assertEqual(values[("Program", "C5")], xl.ERR["#REF!"])
        self.assertEqual(values[("Energy", "B16")], xl.evaluate(xl.workbook())[("Energy", "B16")] - xl.evaluate(xl.workbook())[("Type B", "B13")])
        self.assertEqual(wb.cell("Energy", "B16").formula, "=SUM('Type A:Type D'!B13)")


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.wb = xl.workbook()
        self.values = xl.evaluate(self.wb)

    def test_sheet_order_and_names(self):
        self.assertEqual(self.wb.order(), CASES["workbook"]["sheets"])
        self.assertEqual(set(self.wb.names), {"theta_i", "theta_e", "theta_base", "heating_degree_days", "attached_factor", "air_change", "air_heat_capacity", "u_roof", "u_floor", "u_wall", "u_window"})

    def test_key_results(self):
        for target, expected in CASES["workbook"]["values"].items():
            self.assertTrue(same(self.values[key(target)], expected_value(expected)), (target, self.values[key(target)]))

    def test_no_errors_and_no_cycles(self):
        self.assertEqual([k for k, v in self.values.items() if isinstance(v, xl.XlError)], [])
        self.assertEqual(xl.circular_cells(self.wb), [])

    def test_display_in_both_languages(self):
        for target, texts in CASES["workbook"]["display"].items():
            for language in ("en", "de"):
                self.assertEqual(xl.shown(self.wb, self.values, *key(target), language), texts[language], (target, language))

    def test_formula_bar_in_both_languages(self):
        for target, texts in CASES["workbook"]["formula_bar"].items():
            for language in ("en", "de"):
                self.assertEqual(xl.formula_text(self.wb, *key(target), language), texts[language], (target, language))

    def test_fill_down_reproduces_authored_formulas(self):
        for case in CASES["workbook"]["fill_down"]:
            source = self.wb.cell(case["sheet"], case["source"]).formula
            for target in case["targets"]:
                self.assertEqual(xl.translate(source, case["source"], target), self.wb.cell(case["sheet"], target).formula, (case, target))

    def test_layout_fits_the_screen_block(self):
        layout = CASES["workbook"]["layout"]
        for sheet in self.wb.sheets:
            for address in sheet.cells:
                col, row = xl.split_address(address)
                letter = xl.column_letters(col)
                if sheet.name == "Program" and letter in layout["list_columns"]:
                    continue
                self.assertIn(letter, layout["important_columns"], (sheet.name, address))
                self.assertLessEqual(row, layout["important_rows"], (sheet.name, address))

    def test_every_label_is_bilingual(self):
        for sheet in self.wb.sheets:
            for address, cell in sheet.cells.items():
                if cell.labels:
                    self.assertTrue(cell.labels["en"] and cell.labels["de"], (sheet.name, address))
        self.assertEqual(xl.shown(self.wb, self.values, "Climate", "A5", "en"), "Heating degree days")
        self.assertEqual(xl.shown(self.wb, self.values, "Climate", "A5", "de"), "Heizgradtage")

    def test_precedents_and_dependents_of_the_workbook(self):
        for target, expected in CASES["workbook"]["precedents"].items():
            wanted = set()
            for item in expected:
                sheet, _, rng = item.rpartition("!")
                for row in oxl_cell.cols_from_range(rng):
                    wanted.update((sheet, c) for c in row)
            self.assertEqual(set(xl.precedents(self.wb, key(target))), wanted, target)
        for target, expected in CASES["workbook"]["dependents_of"].items():
            self.assertEqual(sorted(xl.dependents(self.wb, key(target))), sorted(key(e) for e in expected), target)

    def test_what_if_cascades_through_sheets(self):
        base = self.values[("Energy", "D8")]
        events = [{"type": "set-cell", "sheet": "Climate", "cell": "B3", "value": -10}]
        colder = xl.evaluate(xl.apply(self.wb, events))
        self.assertGreater(colder[("Type A", "B13")], self.values[("Type A", "B13")])
        self.assertEqual(colder[("Energy", "D8")], base)

    def test_type_sheets_share_one_layout_for_3d_references(self):
        shape = lambda name: {a: c.formula for a, c in self.wb.sheet(name).cells.items() if a.startswith("B") and c.formula and xl.split_address(a)[1] <= 16}
        for name in ("Type B", "Type C", "Type D"):
            self.assertEqual(shape(name), shape("Type A"), name)
        total = sum(self.values[(f"Type {t}", "B13")] for t in "ABCD")
        self.assertAlmostEqual(self.values[("Energy", "B16")], total)


class OracleWorkbookTests(unittest.TestCase):
    """🧮 The whole Nakagin workbook, written to .xlsx with openpyxl and recalculated by the `formulas` library."""

    @staticmethod
    def expand_3d(formula, wb):
        def swap(match):
            names = wb.order()
            first, last = names.index(match[1]), names.index(match[2])
            return "SUM(" + ",".join("'" + n + "'!" + match[3] for n in names[first:last + 1]) + ")"
        return re.sub(r"SUM\('([^:']+):([^']+)'!(\$?[A-Z]+\$?\d+)\)", swap, formula)

    def test_every_formula_cell_matches_the_formulas_library(self):
        wb = xl.workbook()
        mine = xl.evaluate(wb)
        book = openpyxl.Workbook()
        book.remove(book.active)
        table = program_table()
        for sheet in wb.sheets:
            ws = book.create_sheet(sheet.name)
            for address, cell in sheet.cells.items():
                ws[address] = self.expand_3d(expand_structured(cell.formula, table, xl.split_address(address)[1]), wb) if cell.formula else cell.value
        book["Program"].add_table(openpyxl.worksheet.table.Table(displayName="Capsules", ref="H1:N141"))
        for name, ref in wb.names.items():
            book.defined_names[name] = DefinedName(name, attr_text=ref)
        out = TICKET / "🗑️generated" / "xl-engine"
        out.mkdir(parents=True, exist_ok=True)
        path = out / "Nakagin.xlsx"
        book.save(path)
        reread = openpyxl.load_workbook(path)
        self.assertEqual(reread.sheetnames, wb.order())
        self.assertEqual(reread["Program"].tables["Capsules"].ref, "H1:N141")
        structured = openpyxl.Workbook()
        structured.remove(structured.active)
        for sheet in wb.sheets:
            ws = structured.create_sheet(sheet.name)
            for address, cell in sheet.cells.items():
                ws[address] = cell.formula if cell.formula else cell.value
        structured["Program"].add_table(openpyxl.worksheet.table.Table(displayName="Capsules", ref="H1:N141"))
        structured.save(out / "Nakagin-structured.xlsx")
        again = openpyxl.load_workbook(out / "Nakagin-structured.xlsx")
        self.assertEqual(again["Program"]["B4"].value, "=COUNTIF(Capsules[Type],A4)")
        self.assertEqual(again["Program"].tables["Capsules"].ref, "H1:N141")
        import formulas
        solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
        count = 0
        for (sheet, address), value in mine.items():
            if not wb.cell(sheet, address).formula:
                continue
            got = solution[f"'[Nakagin.xlsx]{sheet.upper()}'!{address}"].value[0, 0]
            if isinstance(value, str):
                self.assertEqual(str(got), value, (sheet, address))
            else:
                self.assertTrue(math.isclose(float(got), float(value), rel_tol=1e-9, abs_tol=1e-9), (sheet, address, got, value))
            count += 1
        self.assertGreater(count, 250)


class GermanOracleWorkbookTests(unittest.TestCase):
    """🇩🇪 The German export: German tab names in every sheet reference, recalculated by `formulas` and equal to our values."""

    def test_german_tab_names_export_matches(self):
        wb = xl.workbook()
        mine = xl.evaluate(xl.workbook(language="de"))
        titles = {s["name"]: s["title"]["de"] for s in xl.load()["sheets"]}
        self.assertEqual(list(titles.values()), ["Programm", "Typ A", "Typ B", "Typ C", "Typ D", "Klima", "Energie"])

        table = program_table()

        def retab(formula, row):
            formula = OracleWorkbookTests.expand_3d(expand_structured(formula, table, row), wb)
            for name, german in titles.items():
                formula = re.sub(r"'" + re.escape(name) + r"'!|(?<![\w'])" + re.escape(name) + r"!", "'" + german + "'!", formula)
            return formula

        book = openpyxl.Workbook()
        book.remove(book.active)
        for sheet in wb.sheets:
            ws = book.create_sheet(titles[sheet.name])
            for address, cell in sheet.cells.items():
                ws[address] = retab(cell.formula, xl.split_address(address)[1]) if cell.formula else (cell.labels["de"] if cell.labels else cell.value)
        for name, ref in wb.names.items():
            sheet, _, cells = ref.partition("!")
            book.defined_names[name] = DefinedName(name, attr_text="'" + titles[sheet] + "'!" + cells)
        out = TICKET / "🗑️generated" / "xl-engine"
        out.mkdir(parents=True, exist_ok=True)
        path = out / "Nakagin-de.xlsx"
        book[titles["Program"]].add_table(openpyxl.worksheet.table.Table(displayName=xl.table_titles("de")["Capsules"], ref="H1:N141"))
        book.save(path)
        self.assertEqual(openpyxl.load_workbook(path).sheetnames, list(titles.values()))
        self.assertEqual(openpyxl.load_workbook(path)[titles["Program"]].tables["Kapseln"].ref, "H1:N141")
        import formulas
        solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
        count = 0
        for (sheet, address), value in mine.items():
            if not wb.cell(sheet, address).formula:
                continue
            got = solution[f"'[Nakagin-de.xlsx]{titles[sheet].upper()}'!{address}"].value[0, 0]
            if isinstance(value, str):
                self.assertEqual(str(got), value)
            else:
                self.assertTrue(math.isclose(float(got), float(value), rel_tol=1e-9, abs_tol=1e-9), (sheet, address, got, value))
            count += 1
        self.assertGreater(count, 250)


class NakaginCrossCheckTests(unittest.TestCase):
    """🏢 The workbook must agree with the shared Nakagin teaching model it is bound to."""

    @classmethod
    def setUpClass(cls):
        if not NAKAGIN.exists():
            raise unittest.SkipTest("Nakagin model is not available")
        cls.nk = load_module("nakagin_model_for_excel", NAKAGIN)
        cls.wb = xl.workbook()
        cls.values = xl.evaluate(cls.wb)

    def test_bound_inputs_equal_the_model_values(self):
        count = 0
        for spec in xl.load()["sheets"]:
            for address, cell in spec["cells"].items():
                if "from" in cell:
                    self.assertEqual(cell["value"], self.nk.value(cell["from"]).value, (spec["name"], address))
                    self.assertEqual(cell["basis"], self.nk.value(cell["from"]).basis, (spec["name"], address))
                    count += 1
        self.assertGreater(count, 80)
        rebound = xl.evaluate(xl.workbook(lambda path: self.nk.value(path).value))
        self.assertEqual(rebound, self.values)

    def test_capsule_list_is_the_slot_list(self):
        slots = self.nk.slots()
        self.assertEqual(len(slots), 140)
        letters = "NESW"
        for index, slot in enumerate(slots):
            row = index + 2
            tower = self.nk.towers()[slot.tower].key
            got = [self.values[("Program", f"{col}{row}")] for col in "HIJKLM"]
            self.assertEqual(got, [f"{tower}-{slot.storey + 1:02d}-{letters[slot.facade]}-{slot.item + 1}", tower, slot.storey + 1, letters[slot.facade], slot.item + 1, f"Type {slot.type}"], row)

    def test_counts_and_areas(self):
        totals = self.nk.totals()
        for index, t in enumerate(self.nk.capsule_types()):
            self.assertEqual(self.values[("Program", f"B{4 + index}")], t.count)
            self.assertAlmostEqual(self.values[("Program", f"C{4 + index}")], self.nk.program(t)["net_floor_area"])
        self.assertEqual(self.values[("Program", "B8")], totals["count"])
        self.assertAlmostEqual(self.values[("Program", "D8")], totals["area"])
        for index, tower in enumerate(self.nk.towers()):
            self.assertEqual(self.values[("Program", f"B{11 + index}")] + self.values[("Program", f"C{11 + index}")] + self.values[("Program", f"D{11 + index}")] + self.values[("Program", f"E{11 + index}")], totals["per_tower"][tower.key]["count"])
            self.assertAlmostEqual(self.values[("Program", f"F{11 + index}")], totals["per_tower"][tower.key]["area"])

    def test_envelope_and_heating_per_type(self):
        for t in self.nk.capsule_types():
            sheet = f"Type {t.key}"
            env, heat, program = self.nk.envelope(t), self.nk.heating(t), self.nk.program(t)
            for row, surface in enumerate(env["surfaces"], start=3):
                self.assertAlmostEqual(self.values[(sheet, f"D{row}")], surface["area"], msg=(sheet, surface["key"]))
                self.assertAlmostEqual(self.values[(sheet, f"E{row}")], surface["u_value"])
                self.assertAlmostEqual(self.values[(sheet, f"F{row}")], surface["ua"])
            self.assertAlmostEqual(self.values[(sheet, "B10")], program["volume"])
            self.assertAlmostEqual(self.values[(sheet, "B11")], env["h_t"])
            self.assertAlmostEqual(self.values[(sheet, "B12")], env["h_v"])
            self.assertAlmostEqual(self.values[(sheet, "B13")], heat["load"])
            self.assertAlmostEqual(self.values[(sheet, "B14")], heat["annual"])
            self.assertAlmostEqual(self.values[(sheet, "B15")], heat["annual_per_area"])
            self.assertAlmostEqual(self.values[(sheet, "G16")], program["fitting_area"])
            self.assertAlmostEqual(self.values[(sheet, "B16")], program["free_area"])

    def test_climate_and_building_totals(self):
        climate, totals = self.nk.climate(), self.nk.totals()
        self.assertAlmostEqual(self.values[("Climate", "B5")], climate["hdd"])
        self.assertAlmostEqual(self.values[("Energy", "D8")], totals["annual"])
        self.assertAlmostEqual(self.values[("Energy", "F8")], totals["annual_per_area"])
        for index, t in enumerate(self.nk.capsule_types()):
            self.assertAlmostEqual(self.values[("Energy", f"F{4 + index}")], self.nk.heating(t)["annual_per_area"])


class StyleTests(unittest.TestCase):
    def test_docstrings_start_with_an_emoji_and_bodies_have_no_comments(self):
        path = CLIP / "🧱️model/🐍️.py"
        source = path.read_text(encoding="utf-8")
        tree = ast.parse(source)
        for node in ast.walk(tree):
            if isinstance(node, (ast.FunctionDef, ast.ClassDef, ast.AsyncFunctionDef)):
                doc = ast.get_docstring(node, clean=False)
                if doc is not None:
                    self.assertGreater(ord(doc.lstrip()[0]), 127, node.name)
        self.assertTrue(ast.get_docstring(tree).lstrip()[0] > "\x7f")
        comments = [t for t in tokenize.generate_tokens(io.StringIO(source).readline) if t.type == tokenize.COMMENT]
        self.assertEqual(comments, [])

    def test_public_api_is_present(self):
        for name in ("Ref", "Range", "parse_ref", "format_ref", "translate", "parse", "evaluate", "precedents", "dependents", "localize", "tab_name", "sheet_titles", "display", "workbook", "apply", "cycle_reference", "recalculation_order", "circular_cells"):
            self.assertTrue(hasattr(xl, name), name)

    def test_no_third_party_runtime_imports(self):
        tree = ast.parse((CLIP / "🧱️model/🐍️.py").read_text(encoding="utf-8"))
        allowed = set(sys.stdlib_module_names)
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                self.assertTrue(all(a.name.split(".")[0] in allowed for a in node.names))
            if isinstance(node, ast.ImportFrom) and node.level == 0:
                self.assertIn(node.module.split(".")[0], allowed)


if __name__ == "__main__":
    unittest.main(verbosity=1)
